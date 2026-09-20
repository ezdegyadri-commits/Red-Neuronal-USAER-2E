from copy import copy
from datetime import date
from io import BytesIO
from pathlib import Path
import re
import unicodedata

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from services.padron_oficial import vista_padron_oficial


TEMPLATES_DIR = Path(__file__).resolve().parents[1] / "assets" / "plantillas"
PADRON_TEMPLATE = "padron_usaer_2026_2027.xlsx"
PERSONAL_TEMPLATE = "personal_usaer_2026_2027.xlsx"

SERVICE = {
    "zona": "001",
    "numero": "02-E",
    "cct": "31FUA0002Y",
    "turno": "Matutino",
}


def _text(value):
    if value is None or pd.isna(value):
        return ""
    return str(value).strip()


def _norm(value):
    return "".join(
        char for char in unicodedata.normalize("NFD", _text(value).upper())
        if unicodedata.category(char) != "Mn"
    )


def _value(row, *keys, default=""):
    for key in keys:
        if key in row and _text(row.get(key)):
            return _text(row.get(key))
    return default


def _yes(value):
    return _norm(value) in {"SI", "S", "YES", "TRUE", "1"}



def _edad_al_primero_de_septiembre(curp):
    curp = _text(curp).upper()
    try:
        year_two_digits = int(curp[4:6])
        nacimiento = date(
            2000 + year_two_digits if year_two_digits <= 26 else 1900 + year_two_digits,
            int(curp[6:8]),
            int(curp[8:10]),
        )
    except (ValueError, IndexError):
        return ""
    corte = date(2026, 9, 1)
    return corte.year - nacimiento.year - (
        (corte.month, corte.day) < (nacimiento.month, nacimiento.day)
    )

def _school_index(escuelas):
    if escuelas is None or escuelas.empty:
        return {}
    resultado = {}
    for _, row in escuelas.fillna("").iterrows():
        registro = row.to_dict()
        for key in ("ID_Escuela", "Nombre_Escuela"):
            valor = _norm(registro.get(key, ""))
            if valor:
                resultado[valor] = registro
    return resultado


def _copy_row_style(worksheet, source_row, target_row):
    for column in range(1, worksheet.max_column + 1):
        source = worksheet.cell(source_row, column)
        target = worksheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)
        if source.fill:
            target.fill = copy(source.fill)
        if source.font:
            target.font = copy(source.font)
        if source.border:
            target.border = copy(source.border)
        if source.protection:
            target.protection = copy(source.protection)
    worksheet.row_dimensions[target_row].height = (
        worksheet.row_dimensions[source_row].height
    )


def _prepare_rows(
    worksheet, first_row, total_rows, source_row, template_last_row
):
    required_last = max(
        template_last_row, first_row + max(total_rows, 1) - 1
    )
    for row in range(template_last_row + 1, required_last + 1):
        _copy_row_style(worksheet, source_row, row)
    for row in range(first_row, required_last + 1):
        for column in range(1, worksheet.max_column + 1):
            cell = worksheet.cell(row, column)
            if cell.data_type != "f":
                cell.value = None
            cell.hyperlink = None


def _style_header(cell, fill):
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(name="Arial", size=8, bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(
        left=Side(style="thin", color="666666"),
        right=Side(style="thin", color="666666"),
        top=Side(style="thin", color="666666"),
        bottom=Side(style="thin", color="666666"),
    )


def _crear_padron_base():
    workbook = Workbook()
    instrucciones = workbook.active
    instrucciones.title = "INSTRUCTIVO"
    instrucciones["A1"] = "PADRÓN DE ALUMNOS ATENDIDOS EN EL CURSO ESCOLAR 2026-2027"
    instrucciones["A3"] = "El nombre inicia con apellido paterno, materno y nombre(s)."
    instrucciones["A4"] = "La CURP debe contener 18 caracteres."
    worksheet = workbook.create_sheet("ALUMNOS USAER")
    worksheet.merge_cells("A3:Y3")
    worksheet["A3"] = "PADRÓN DE ALUMNOS ATENDIDOS POR LA USAER CURSO ESCOLAR 2026-2027"
    worksheet["A3"].font = Font(name="Arial", size=12, bold=True)
    worksheet["A3"].alignment = Alignment(horizontal="center")
    headers = [
        "No.", "Zona", "N° USAER", "Clave de USAER", "Escuela atendida",
        "Turno", "CCT de la escuela", "Dirección de la escuela", "Localidad",
        "Municipio", "Apellido paterno, materno y nombre(s) del alumno",
        "CURP (18 dígitos)", "Edad (1 sept)", "Sexo", "Discapacidad o condición",
        "Preescolar", "Primaria", "Secundaria", "Situación del alumno",
        "Tipo de atención", "Lengua indígena / mayahablante",
        "Afrodescendiente", "Migrante", "", "",
    ]
    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(6, column)
        cell.value = header
        _style_header(cell, "B4A7C5")
        worksheet.column_dimensions[cell.column_letter].width = 15
        worksheet.cell(8, column).border = Border(
            left=Side(style="thin", color="999999"),
            right=Side(style="thin", color="999999"),
            top=Side(style="thin", color="999999"),
            bottom=Side(style="thin", color="999999"),
        )
    worksheet.row_dimensions[6].height = 58
    worksheet.freeze_panes = "A7"
    return workbook


def _crear_personal_base():
    workbook = Workbook()
    instrucciones = workbook.active
    instrucciones.title = "INSTRUCCIONES"
    instrucciones["A1"] = "FORMATO DE PERSONAL DE USAER 2026-2027"
    worksheet = workbook.create_sheet("1. SÁBANA DE PERSONAL DE USAER")
    worksheet.merge_cells("A3:CP3")
    worksheet["A3"] = "SÁBANA DE PERSONAL DE USAER"
    worksheet["A3"].font = Font(name="Arial", size=12, bold=True)
    worksheet["A3"].alignment = Alignment(horizontal="center")
    worksheet["A5"] = "Fecha de llenado:"
    headers = [
        "No.", "Zona", "N° USAER", "Clave de USAER", "Turno",
        "Nombre del docente de apoyo/paradocentes", "Sexo", "Función",
        "Correo electrónico", "Teléfono", "Base/Contrato", "Sostenimiento",
        "¿Presenta alguna discapacidad?", "Horario de trabajo", "¿Cuál?",
        "¿Es maya hablante?", "Número de escuela", "Nombre de la escuela que atiende",
        "CCT de la escuela regular", "Nivel educativo", "Modalidad",
        "Horario de la escuela", "Total de grupos", "Dirección de la escuela",
        "Localidad", "Municipio", "Nombre del director(a)", "Teléfono del director",
        "Nombre del supervisor(a)", "Teléfono del supervisor", "Zona escolar",
        "Sector", "Región", "Alumnos grupales", "Alumnos individuales",
        "Total alumnos", "EPP vigentes", "EPP anteriores", "Total EPP",
        "Ceguera H", "Ceguera M", "Baja visión H", "Baja visión M", "Sordera H",
        "Sordera M", "Hipoacusia H", "Hipoacusia M", "Sordoceguera H",
        "Sordoceguera M", "Motriz H", "Motriz M", "Intelectual H", "Intelectual M",
        "Psicosocial H", "Psicosocial M", "TEA H", "TEA M", "Múltiple H",
        "Múltiple M", "TDA H", "TDA M", "TDAH H", "TDAH M", "AS intelectual H",
        "AS intelectual M", "AS artístico H", "AS artístico M", "AS psicomotriz H",
        "AS psicomotriz M", "AS socioafectiva H", "AS socioafectiva M",
        "AS creativa H", "AS creativa M", "Aprendizaje H", "Aprendizaje M",
        "Conducta H", "Conducta M", "Comunicación H", "Comunicación M",
        "Otras condiciones H", "Otras condiciones M", "Otras total",
        "Total alumnos H", "Total alumnos M", "Total alumnos",
        "Maya hablantes H", "Maya hablantes M", "Afrodescendientes H",
        "Afrodescendientes M", "Migrantes H", "Migrantes M",
        "Doble vulnerabilidad H", "Doble vulnerabilidad M",
        "Doble vulnerabilidad total",
    ]
    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(8, column)
        cell.value = header
        fill = "9DC3E6" if column <= 5 else "B6D7A8" if column <= 16 else "F8CBAD"
        _style_header(cell, fill)
        worksheet.column_dimensions[cell.column_letter].width = 15
        worksheet.cell(10, column).border = Border(
            left=Side(style="thin", color="999999"),
            right=Side(style="thin", color="999999"),
            top=Side(style="thin", color="999999"),
            bottom=Side(style="thin", color="999999"),
        )
    worksheet.row_dimensions[8].height = 72
    worksheet.freeze_panes = "A9"
    return workbook


def _workbook_from_template(filename):
    path = TEMPLATES_DIR / filename
    if path.exists():
        return load_workbook(BytesIO(path.read_bytes()))
    if filename == PADRON_TEMPLATE:
        return _crear_padron_base()
    if filename == PERSONAL_TEMPLATE:
        return _crear_personal_base()
    raise FileNotFoundError(f"No se encontró la plantilla oficial: {filename}")


def _to_excel(workbook):
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def generar_padron_usaer(alumnos, escuelas):
    """Genera el padrón USAER en la plantilla oficial de supervisión."""
    workbook = _workbook_from_template(PADRON_TEMPLATE)
    worksheet = workbook["ALUMNOS USAER"]
    registros = vista_padron_oficial(alumnos, escuelas).fillna("").to_dict("records")
    first_row = 8
    _prepare_rows(worksheet, first_row, len(registros), first_row, 30)

    for number, student in enumerate(registros, start=1):
        row = first_row + number - 1
        values = {
            1: number,
            2: student["ZONA"], 3: student["USAER"], 4: student["CLAVE DE USAER"],
            5: student["ESCUELA ATENDIDA"], 6: student["TURNO"],
            7: student["CCT DE LA ESCUELA"], 8: student["DIRECCIÓN DE LA ESCUELA"],
            9: student["LOCALIDAD"], 10: student["MUNICIPIO"],
            11: student["NOMBRE COMPLETO DEL ALUMNO"], 12: student["CURP"],
            13: student["EDAD AL 1 DE SEPTIEMBRE"], 14: student["SEXO"],
            15: student["DISCAPACIDAD O CONDICIÓN"], 16: student["PREESCOLAR"],
            17: student["PRIMARIA"], 18: student["SECUNDARIA"],
            19: student["SITUACIÓN DEL ALUMNO"], 20: student["TIPO DE ATENCIÓN"],
            21: student["MAYA HABLANTE"], 22: student["MIGRANTE"],
            23: student["AFRODESCENDIENTE"],
        }
        for column, value in values.items():
            worksheet.cell(row, column).value = _text(value)
        worksheet.cell(row, 2).number_format = "@"
    return _to_excel(workbook)


def _students_for_school(alumnos, school_id):
    if alumnos is None or alumnos.empty:
        return []
    target = _norm(school_id)
    return [
        row for row in alumnos.fillna("").to_dict("records")
        if target and target in {
            _norm(row.get("ID_Escuela", "")),
            _norm(row.get("CCT_Escuela", "")),
            _norm(row.get("Nombre_Escuela", "")),
        }
    ]


def _sex_count(records, predicate=lambda record: True):
    selected = [row for row in records if predicate(row)]
    men = sum(_norm(row.get("Sexo", "")) in {"H", "HOMBRE", "MASCULINO"} for row in selected)
    women = sum(_norm(row.get("Sexo", "")) in {"M", "MUJER", "FEMENINO"} for row in selected)
    return men, women


def _write_condition_counts(worksheet, row, students):
    columns = {
        "CEGUERA": 40,
        "BAJA VISION": 42,
        "SORDERA": 44,
        "HIPOACUSIA": 46,
        "SORDOCEGUERA": 48,
        "MOTORA": 50,
        "INTELECTUAL": 52,
        "PSICOSOCIAL": 54,
        "TEA": 56,
        "MULTIPLE": 58,
        "TDA": 60,
        "TDAH": 62,
        "AS": 64,
        "APRENDIZAJE": 74,
        "CONDUCTA": 76,
        "COMUNICACION": 78,
    }
    for condition, column in columns.items():
        def matches(student):
            recorded = _norm(student.get("Condicion_Discapacidad", ""))
            if condition == "TDA":
                return recorded == "TDA"
            if condition == "TDAH":
                return "TDAH" in recorded
            return condition in recorded
        men, women = _sex_count(students, matches)
        worksheet.cell(row, column).value = men
        worksheet.cell(row, column + 1).value = women


_SCHOOL_ORDER = {
    "DAMIAN CARMONA": 1,
    "ICHCAANZIHO": 2,
    "GREGORIO TORRES QUINTERO": 3,
    "REMIGIO AGUILAR SOSA": 4,
    "ELVIRA PARRA AVILA": 5,
    "MANUEL SARRADO": 6,
    "DOMINGO SOLIS RODRIGUEZ": 7,
    "QUINTANA ROO": 8,
}


def _school_names(value):
    text = _text(value).strip().strip('"')
    if not text:
        return []
    parts = re.split(r"\s*(?:;|,|\by\b)\s*", text, flags=re.IGNORECASE)
    canonical = {
        "DAMIAN CAMONA": "Damián Carmona",
        "DAMIAN CARMONA": "Damián Carmona",
        "ICHC AANZIHO": "Ichcaanzihó",
        "ICHCAANZIHO": "Ichcaanzihó",
        "IHCAANZIHO": "Ichcaanzihó",
        "IHC AANZIHO": "Ichcaanzihó",
        "GREGORIO TORRES QUINTERO": "Gregorio Torres Quintero",
        "REMIGIO AGUILAR SOSA": "Remigio Aguilar Sosa",
        "ELVIRA PARRA AVILA": "Elvira Parra Ávila",
        "MANUEL SARRADO": "Manuel Sarrado",
        "DOMINGO SOLIS": "Domingo Solís Rodríguez",
        "DOMINGO SOLIS RODRIGUEZ": "Domingo Solís Rodríguez",
        "QUINTANA RO0": "Quintana Roo",
        "QUINTANA ROO": "Quintana Roo",
    }
    result = []
    seen = set()
    for part in parts:
        name = part.strip().strip('"').rstrip(".")
        if not name:
            continue
        normalized = _norm(name)
        name = canonical.get(normalized, name)
        key = _norm(name)
        if key not in seen:
            result.append(name)
            seen.add(key)
    return result


def _ordered_schools(names):
    return sorted(
        names,
        key=lambda name: (
            _SCHOOL_ORDER.get(_norm(name), 99),
            _norm(name),
        ),
    )


def _role_order(person):
    role = _norm(person.get("Rol", ""))
    if "DIRECTOR" in role:
        return 0
    if "ADMINISTRATIVO" in role:
        return 1
    if "MAESTRO DE APOYO" in role:
        return 2
    if "COMUNICACION" in role:
        return 3
    if "PSICOLOG" in role:
        return 4
    if "SOCIAL" in role:
        return 5
    return 6


def _assignment_rows(personal, asignaciones):
    people = [] if personal is None else personal.fillna("").to_dict("records")
    assignments = [] if asignaciones is None else asignaciones.fillna("").to_dict("records")
    by_person = {}
    for item in assignments:
        by_person.setdefault(_norm(item.get("ID_Personal", "")), []).append(
            _text(item.get("ID_Escuela", ""))
        )
    rows = []
    for person in people:
        if person.get("Fuente_Formulario"):
            schools_for_person = _school_names(person.get("Escuelas_Atendidas", ""))
            if _role_order(person) == 2 and len(schools_for_person) > 1:
                for school_name in _ordered_schools(schools_for_person):
                    person_row = dict(person)
                    person_row["Escuela_Asignada"] = school_name
                    for field in (
                        "CCT_Escuela", "Nivel_Escuela", "Modalidad_Escuela",
                        "Horario_Escuela", "Grupos_Escuela", "Direccion_Escuela",
                        "Localidad_Escuela", "Municipio_Escuela",
                    ):
                        person_row[field] = ""
                    rows.append((person_row, school_name))
            else:
                rows.append((person, _text(person.get("Escuela_Asignada", ""))))
            continue
        assigned = by_person.get(_norm(person.get("ID_Personal", "")), [""])
        for school_id in assigned or [""]:
            rows.append((person, school_id))
    if any(person.get("Fuente_Formulario") for person, _ in rows):
        rows.sort(
            key=lambda item: (
                _role_order(item[0]),
                _SCHOOL_ORDER.get(_norm(item[1]), 99),
                _norm(item[0].get("Nombre_Completo", "")),
            )
        )
    return rows


def _students_for_person(alumnos, person, school, school_id):
    if not person.get("Fuente_Formulario"):
        return _students_for_school(alumnos, school_id)
    if "ADMINISTRATIVO" in _norm(person.get("Rol", "")):
        return []
    names = _school_names(person.get("Escuelas_Atendidas", ""))
    assigned_name = _text(person.get("Escuela_Asignada", ""))
    if assigned_name:
        names = [assigned_name]
    keys = {_norm(value) for value in names if _text(value)}
    cct = _text(person.get("CCT_Escuela", ""))
    if cct:
        keys.add(_norm(cct))
    if not keys or alumnos is None or alumnos.empty:
        return []
    selected = []
    seen = set()
    for student in alumnos.fillna("").to_dict("records"):
        if not keys.intersection({
            _norm(student.get("ID_Escuela", "")),
            _norm(student.get("CCT_Escuela", "")),
            _norm(student.get("Nombre_Escuela", "")),
        }):
            continue
        unique = _norm(student.get("ID_Alumno", "") or student.get("CURP", ""))
        unique = unique or str(len(selected))
        if unique not in seen:
            seen.add(unique)
            selected.append(student)
    return selected


def generar_formato_personal(personal, escuelas, alumnos, asignaciones):
    """Genera la sábana de personal usando la plantilla de supervisión."""
    workbook = _workbook_from_template(PERSONAL_TEMPLATE)
    worksheet = workbook["1. SÁBANA DE PERSONAL DE USAER"]
    worksheet["D5"] = date.today().strftime("%d/%m/%Y")
    rows = _assignment_rows(personal, asignaciones)
    schools = _school_index(escuelas)
    first_row = 10
    _prepare_rows(worksheet, first_row, len(rows), first_row, 60)

    for number, (person, school_id) in enumerate(rows, start=1):
        row = first_row + number - 1
        formulario = bool(person.get("Fuente_Formulario"))
        school_names = _school_names(person.get("Escuelas_Atendidas", "")) if formulario else []
        role = _norm(person.get("Rol", ""))
        if formulario and "ADMINISTRATIVO" in role:
            school = {}
            school_number = "SEDE"
            school_name = "USAER 02E"
        elif formulario and _role_order(person) >= 3:
            school = {}
            ordered = _ordered_schools(school_names)
            school_number = ", ".join(
                str(_SCHOOL_ORDER.get(_norm(name), "")) for name in ordered
            ).strip(", ")
            school_name = "; ".join(ordered)
        elif formulario:
            school = person
            assigned = _text(person.get("Escuela_Asignada", "")) or (
                school_names[0] if school_names else ""
            )
            school_number = _SCHOOL_ORDER.get(_norm(assigned), "")
            school_name = assigned
        else:
            school = schools.get(_norm(school_id), {})
            school_number = _value(school, "Numero_Escuela")
            school_name = _value(school, "Nombre_Escuela")
        students = _students_for_person(alumnos, person, school, school_id)
        group_count = sum(
            _norm(student.get("Tipo_Atencion", "")) == "GRUPAL"
            for student in students
        )
        individual_count = sum(
            _norm(student.get("Tipo_Atencion", "")) == "INDIVIDUAL"
            for student in students
        )
        values = {
            1: number,
            2: _value(person, "Zona", default=SERVICE["zona"]),
            3: _value(person, "Numero_USAER", default=SERVICE["numero"]),
            4: _value(person, "CCT_USAER", default=SERVICE["cct"]).replace("-", "").replace(" ", ""),
            5: _value(person, "Turno_USAER", default=SERVICE["turno"]),
            6: _value(person, "Nombre_Completo", "Nombre"),
            7: _value(person, "Sexo"),
            8: _value(person, "Rol", "Funcion"),
            9: _value(person, "Email", "Correo_Electronico"),
            10: _value(person, "Telefono", "Teléfono"),
            11: _value(person, "Base_Contrato"),
            12: _value(person, "Sostenimiento"),
            13: _value(person, "Presenta_Discapacidad", default="No"),
            14: _value(person, "Horario"),
            15: _value(person, "Discapacidad"),
            16: _value(person, "Maya_Hablante", default="No"),
            17: school_number,
            18: school_name,
            19: _value(school, "CCT", "CCT_Escuela"),
            20: _value(school, "Nivel", "Nivel_Escuela"),
            21: _value(school, "Modalidad", "Modalidad_Escuela"),
            22: _value(school, "Horario", "Horario_Escuela"),
            23: _value(school, "Total_Grupos", "Grupos_Escuela"),
            24: _value(school, "Direccion", "Dirección", "Direccion_Escuela"),
            25: _value(school, "Localidad", "Localidad_Escuela"),
            26: _value(school, "Municipio", "Municipio_Escuela"),
            27: _value(school, "Director"),
            28: _value(school, "Telefono_Director"),
            29: _value(school, "Supervisor"),
            30: _value(school, "Telefono_Supervisor"),
            31: _value(school, "Zona_Escolar"),
            32: _value(school, "Sector"),
            33: _value(school, "Region"),
            34: group_count,
            35: individual_count,
            36: len(students),
            37: individual_count,
            38: 0,
            39: individual_count,
        }
        for column, value in values.items():
            worksheet.cell(row, column).value = value
        worksheet.cell(row, 2).number_format = "@"
        if formulario and _role_order(person) >= 3:
            for column in (17, 18):
                cell = worksheet.cell(row, column)
                cell.alignment = copy(cell.alignment)
                cell.alignment = cell.alignment.copy(wrap_text=True, vertical="center")
            worksheet.row_dimensions[row].height = max(
                worksheet.row_dimensions[row].height or 0, 72
            )

        _write_condition_counts(worksheet, row, students)
        men, women = _sex_count(students)
        worksheet.cell(row, 83).value = men
        worksheet.cell(row, 84).value = women
        worksheet.cell(row, 85).value = len(students)

        for condition, column in (
            ("Lengua_Indigena_Mayahablante", 86),
            ("Afrodescendiente", 88),
            ("Migrante", 90),
        ):
            men, women = _sex_count(
                students, lambda student, field=condition: _yes(student.get(field, ""))
            )
            worksheet.cell(row, column).value = men
            worksheet.cell(row, column + 1).value = women
        men, women = _sex_count(
            students,
            lambda student: any(
                _yes(student.get(field, ""))
                for field in (
                    "Lengua_Indigena_Mayahablante", "Afrodescendiente", "Migrante"
                )
            ),
        )
        worksheet.cell(row, 92).value = men
        worksheet.cell(row, 93).value = women
        worksheet.cell(row, 94).value = men + women
    return _to_excel(workbook)

