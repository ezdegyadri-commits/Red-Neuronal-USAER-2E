"""Conciliación explícita por CURP. Ausencia en Excel nunca significa baja."""
from copy import deepcopy
from datetime import datetime
from io import BytesIO
import hashlib
import json
import re
import threading
from zoneinfo import ZoneInfo
from utils.text import normalizar_texto
from utils.escolar import grado_canonico, grupo_canonico

_lock=threading.RLock()
ALIASES={
 'Nombre_Completo':('NOMBRE COMPLETO','NOMBRE COMPLETO DEL ALUMNO','APELLIDO PATERNO MATERNO Y NOMBRE S COMPLETO DEL ALUMNO','APELLIDO PATERNO MATERNO Y NOMBRE S COMPLETO'),
 'CURP':('CURP','CURP 18 DIGITOS'), 'Grado':('GRADO','NIVEL Y GRADO','NIVEL Y GRADO AL QUE ESTA INSCRITO'),
 'Grupo':('GRUPO',), 'Edad_1_Septiembre':('EDAD','EDAD 30 SEPT','EDAD 1 SEPT','EDAD 1 SEPTIEMBRE','EDAD AL 1 DE SEPTIEMBRE DE 2026'),
 'Sexo':('SEXO',),'Situacion_Alumno':('SITUACION DEL ALUMNO','SITUACION ALUMNO'),
 'Estatus':('ESTATUS',),
 'Tipo_Atencion':('TIPO DE ATENCION','TIPO ATENCION'),
 'Condicion_Discapacidad':('DISCAPACIDAD O CONDICION','CONDICION DISCAPACIDAD','CONDICION'),
 'Nivel_Educativo':('NIVEL EDUCATIVO',), 'ID_Escuela':('ID ESCUELA',),
 'Nombre_Escuela':('NOMBRE ESCUELA','ESCUELA','NOMBRE DE LA ESCUELA PREESCOLAR PRIMARIA Y O SECUNDARIA ATENDIDA'),
 'CCT_Escuela':('CCT','CCT DE LA ESCUELA','CCT DE LA ESCUELA ATENDIDA'),
 'Turno_Escuela':('TURNO',),'Direccion_Escuela':('DIRECCION','DIRECCION DE LA ESCUELA ATENDIDA'),
 'Localidad_Escuela':('LOCALIDAD','LOCALIDAD DONDE ESTA UBICADA LA ESCUELA'),
 'Municipio_Escuela':('MUNICIPIO','MUNICIPIO DONDE ESTA UBICADA LA ESCUELA'),
 'Lengua_Indigena_Mayahablante':('LENGUA INDIGENA','LENGUA INDIGENA MAYAHABLANTE'),
 'Afrodescendiente':('AFRODESCENDIENTE',),'Migrante':('MIGRANTE',),
 'Fecha_Nacimiento':('FECHA NACIMIENTO','FECHA DE NACIMIENTO'),
 'Fecha_Derivacion':('FECHA DERIVACION','FECHA DE DERIVACION'),
 'ID_Maestro_Regular':('ID MAESTRO REGULAR','DOCENTE REGULAR'),
 'Condiciones_Adicionales':('CONDICIONES ADICIONALES',),
}

def normalizar_header(value):
    return re.sub(r'[^A-Z0-9]+',' ',re.sub(r'^\s*\d+\s*[.\-:]*\s*','',normalizar_texto(str(value or '')).replace('_',' '))).strip()

def texto(value):
    return str(value if value is not None else '').strip()

def leer(data, escuelas):
    from openpyxl import load_workbook
    if len(data)>16*1024*1024:raise ValueError('Usa un Excel .xlsx de hasta 16 MB.')
    from services.materiales_planeacion import _zip
    with _zip(data):pass
    libro=load_workbook(BytesIO(data),read_only=True,data_only=True)
    try:
        candidatos=[s for s in libro if 'CAM' not in normalizar_texto(s.title)]
        candidatos.sort(key=lambda s:0 if 'USAER' in normalizar_texto(s.title) else 1)
        for hoja in candidatos:
            if hoja.max_row>10000 or hoja.max_column>100:raise ValueError('El libro excede el límite de 10000 filas y 100 columnas por hoja.')
            rows=list(hoja.iter_rows(values_only=True))
            cabecera=next((i for i,r in enumerate(rows[:80]) if any(normalizar_header(c).startswith('CURP') for c in r) and any('NOMBRE' in normalizar_header(c) for c in r)),None)
            if cabecera is None:continue
            indices={}
            for i,v in enumerate(rows[cabecera]):
                key=normalizar_header(v)
                for campo,nombres in ALIASES.items():
                    if key in (*nombres,normalizar_header(campo)):indices[campo]=i
            # La plantilla oficial usa dos filas de encabezados y la columna
            # PRIMARIA contiene el grado; Grupo puede estar repetido al final.
            if cabecera+1<len(rows):
                for i,v in enumerate(rows[cabecera+1]):
                    if normalizar_header(v) in ('PREESCOLAR','PRIMARIA','SECUNDARIA'):
                        indices.setdefault('_niveles',[]).append((i,normalizar_header(v).title()))
                    for campo in ('Grupo','Lengua_Indigena_Mayahablante','Afrodescendiente','Migrante'):
                        if normalizar_header(v) in ALIASES[campo] and campo not in indices:indices[campo]=i
            if not {'Nombre_Completo','CURP'}<=set(indices):raise ValueError('No se reconocen nombre y CURP. Usa la plantilla institucional o encabezados explícitos.')
            result=[];errores=[];meta={}
            school_keys={normalizar_header(e.get(k)):e for e in escuelas for k in ('ID_Escuela','Nombre_Escuela','CCT_Escuela') if e.get(k)}
            for n,row in enumerate(rows[cabecera+1:],cabecera+2):
                def val(i):return row[i] if i<len(row) else None
                if not texto(val(indices['CURP'])) or normalizar_header(val(indices['CURP'])).startswith('CURP'):continue
                reg={c:texto(val(i)) for c,i in indices.items() if c!='_niveles' and texto(val(i))}
                for c in ('ID_Escuela','Nombre_Escuela','CCT_Escuela','Turno_Escuela','Direccion_Escuela','Localidad_Escuela','Municipio_Escuela'):
                    if reg.get(c):meta[c]=reg[c]
                    elif c in meta:reg[c]=meta[c]
                for i,nivel in indices.get('_niveles',[]):
                    if re.search(r'[1-6]',texto(val(i))):reg.update(Nivel_Educativo=nivel,Grado=texto(val(i)))
                curp=reg['CURP'].upper();reg['CURP']=curp
                if not re.fullmatch(r'[A-Z]{4}\d{6}[HM][A-Z]{5}[A-Z0-9]\d',curp):errores.append(f'Fila {n}: CURP con formato inválido.');continue
                opciones=[school_keys.get(normalizar_header(reg[c])) for c in ('ID_Escuela','Nombre_Escuela','CCT_Escuela') if reg.get(c)]
                if not opciones and len(escuelas)==1:opciones=[escuelas[0]]
                if not opciones or any(e is None for e in opciones) or len({e['ID_Escuela'] for e in opciones})!=1:
                    errores.append(f'Fila {n}: escuela o CCT no coincide exactamente con tu asignación.');continue
                e=opciones[0]
                for c in ('ID_Escuela','Nombre_Escuela','CCT_Escuela'):reg[c]=e.get(c,reg.get(c,''))
                from utils.escolar import separar_grado_grupo
                try:
                    grado,grupo=separar_grado_grupo(reg.get('Grado',''),reg.get('Grupo',''))
                    if grado:reg['Grado']=grado
                    if grupo:reg['Grupo']=grupo
                except ValueError:errores.append(f'Fila {n}: grado y grupo contradictorios.');continue
                from services.padron_catalogos import canonico,fecha_documental
                for c in ('Tipo_Atencion','Nivel_Educativo','Situacion_Alumno','Condicion_Discapacidad','Lengua_Indigena_Mayahablante','Afrodescendiente','Migrante','Estatus'):
                    if c in reg:reg[c]=canonico(c,reg[c])
                if reg.get('Sexo'):reg['Sexo']={'HOMBRE':'H','MUJER':'M'}.get(normalizar_texto(reg['Sexo']),reg['Sexo'].upper())
                try:
                    if reg.get('Edad_1_Septiembre'):
                        edad=float(reg['Edad_1_Septiembre'])
                        if not edad.is_integer() or not 0<=edad<=100:raise ValueError()
                        reg['Edad_1_Septiembre']=str(int(edad))
                    for c in ('Fecha_Nacimiento','Fecha_Derivacion'):
                        if c in reg:reg[c]=fecha_documental(reg[c]).isoformat()
                except ValueError:errores.append(f'Fila {n}: revisa edad o fechas.');continue
                result.append(reg)
            reconocidas={i for c,i in indices.items() if c!='_niveles'}|{i for i,_ in indices.get('_niveles',[])}
            omitidas=[str(v) for i,v in enumerate(rows[cabecera]) if v and i not in reconocidas]
            return {'registros':result,'errores':errores,'hoja':hoja.title,'sha256':hashlib.sha256(data).hexdigest(),
                'columnas_reconocidas':[c for c in indices if c!='_niveles'],'encabezados_no_reconocidos':omitidas}
        raise ValueError('No se encontró una hoja con encabezados de nombre y CURP.')
    finally:libro.close()

def comparar(registros, actuales, escuelas):
    from services.padron_catalogos import CATALOGOS
    permitidas={str(e['ID_Escuela']) for e in escuelas};index={};seen=set();salida=[];conflictos=[]
    for r in actuales:
        if texto(r.get('Estatus')).upper().startswith('DUPLICADO'):continue
        index.setdefault(texto(r.get('CURP')).upper(),[]).append(r)
    for n,r in enumerate(registros,1):
        invalidos=[k for k,v in r.items() if v and k in CATALOGOS and v not in CATALOGOS[k]]
        if invalidos:
            conflictos.append(f'Registro {n}: categorías fuera del catálogo: '+', '.join(invalidos)+'. Corrige el archivo antes de importar.');continue
        curp=r['CURP'];hits=index.get(curp,[])
        if curp in seen or len(hits)>1:conflictos.append(f'Registro {n}: CURP repetida en el archivo o en la base.');continue
        seen.add(curp)
        if r['ID_Escuela'] not in permitidas or hits and texto(hits[0].get('ID_Escuela'))!=r['ID_Escuela']:
            conflictos.append(f'Registro {n}: cambio de escuela requiere el trámite de traslado.');continue
        viejo=hits[0] if hits else {}
        cambios={k:v for k,v in r.items() if v!='' and texto(v)!=texto(viejo.get(k))}
        if not hits:
            faltan=[c for c in ('Nombre_Completo','Grado','Grupo','Sexo','Tipo_Atencion') if not r.get(c)]
            if faltan:conflictos.append(f'Registro {n}: faltan '+', '.join(faltan));continue
        salida.append({'tipo':'Alta' if not hits else 'Actualizar' if cambios else 'Sin cambios',
            'registro':deepcopy(r),'anterior':deepcopy(viejo),'cambios':cambios})
    return {'filas':salida,'conflictos':conflictos}

def confirmar(plan, escuelas, nombre, rol, archivo):
    from services.asignaciones import puede_dar_alta_alumnos,escuelas_asignadas
    from data import repository as repo
    if not puede_dar_alta_alumnos(nombre,rol):raise PermissionError('Esta cuenta no puede importar el padrón.')
    nombres={normalizar_texto(n) for n in escuelas_asignadas(nombre,rol)}
    if any(normalizar_texto(e['Nombre_Escuela']) not in nombres for e in escuelas):raise PermissionError('Escuela fuera de la asignación vigente.')
    with _lock:
        ws=repo.worksheet('Alumnos');values=repo.retry_google(ws.get_all_values);headers=values[0]
        actuales=[dict(zip(headers,r)) for r in values[1:] if r and r[0]]
        ids=[texto(r.get('ID_Alumno')) for r in actuales]
        if len(set(ids))!=len(ids):raise ValueError('Hay identificadores de alumno repetidos. Dirección debe conciliarlos antes de importar; no se modificó el padrón.')
        nuevo=comparar([r['registro'] for r in plan['filas']],actuales,escuelas)
        if nuevo['conflictos']:raise ValueError('Hay conflictos de identidad o escuela; vuelve a revisar la vista previa.')
        if len(nuevo['filas'])!=len(plan['filas']):raise ValueError('Cambió el padrón; revisa la vista previa.')
        if any(any(texto(a['anterior'].get(k))!=texto(b['anterior'].get(k)) for k in set(a['registro'])|{'ID_Alumno','ID_Escuela','Estatus'}) for a,b in zip(plan['filas'],nuevo['filas'])):
            # Un reintento cuyo resultado ya coincide no escribe otra vez.
            if all(r['tipo']=='Sin cambios' for r in nuevo['filas']):return {'altas':0,'actualizados':0}
            raise ValueError('Cambió el padrón desde la vista previa. Revisa nuevamente antes de confirmar.')
        indices={texto(r.get('ID_Alumno')):i+2 for i,r in enumerate([dict(zip(headers,r)) for r in values[1:]])}
        numero=max([int(texto(r.get('ID_Alumno')).split('-')[-1]) for r in actuales if re.fullmatch(r'ALU-\d+',texto(r.get('ID_Alumno')))]+[0])+1
        requests=[];altas=[];hist=[];actualizados=0
        ahora=datetime.now(ZoneInfo('America/Mexico_City')).isoformat()
        def celda(v):return {'userEnteredValue':{'stringValue':str(v)}}
        for row in nuevo['filas']:
            if row['tipo']=='Sin cambios':continue
            reg=row['registro'];old=row['anterior']
            aid=old.get('ID_Alumno')
            if row['tipo']=='Alta':
                desconocidos=set(reg)-set(headers)
                if desconocidos:raise ValueError('Faltan columnas necesarias en la base central: '+', '.join(sorted(desconocidos)))
                aid=f'ALU-{numero:03d}';numero+=1
                altas.append({**reg,'ID_Alumno':aid,'Estatus':'Activo','Maestra de Apoyo':nombre})
            else:
                for k,v in row['cambios'].items():
                    if k not in headers:raise ValueError('Falta una columna necesaria en la base central: '+k)
                    requests.append({'updateCells':{'start':{'sheetId':ws.id,'rowIndex':indices[aid]-1,'columnIndex':headers.index(k)},'rows':[{'values':[celda(v)]}],'fields':'userEnteredValue'}})
                actualizados+=1
            hist.append({'values':[celda(v) for v in (aid,'2026-2027','IMPORTACION_EXCEL',json.dumps(old,ensure_ascii=False),ahora,'Archivo '+archivo+'; cambios: '+json.dumps(row['cambios'],ensure_ascii=False))]})
        if altas:requests.append({'appendCells':{'sheetId':ws.id,'rows':[{'values':[celda(r.get(k,'')) for k in headers]} for r in altas],'fields':'userEnteredValue'}})
        if hist:
            historial=repo.worksheet('Historial_Correcciones_Padron')
            requests.append({'appendCells':{'sheetId':historial.id,'rows':hist,'fields':'userEnteredValue'}})
        if requests:
            # Una sola transacción literal, incluidos el historial y las altas.
            # No reintentar a ciegas un append cuyo acuse pudo perderse.
            ws.spreadsheet.batch_update({'requests':requests})
            repo.clear_cache('Alumnos');repo.clear_alumnos_cache()
        return {'altas':len(altas),'actualizados':actualizados}
