"""Lectura local de materiales. No ejecuta archivos ni llama a proveedores de IA."""
from io import BytesIO
from pathlib import Path
import hashlib
import re
import zipfile
from xml.etree import ElementTree as ET

TIPOS = ['docx','pdf','png','jpg','jpeg','webp','bmp','tif','tiff','txt','md','csv','json','rtf','odt','xlsx','pptx','odp']
MAX_BYTES = 32 * 1024 * 1024
MAX_TEXTO = 100000
MAX_PAGINAS = 1000
MAX_OCR = 10
MAX_TOTAL_TEXTO = 120000

def _zip(data):
    archive=zipfile.ZipFile(BytesIO(data))
    if len(archive.infolist())>5000 or sum(i.file_size for i in archive.infolist())>128*1024*1024:
        archive.close()
        raise ValueError('El documento es demasiado grande al abrirlo. Divide el material en partes.')
    return archive

def _xml(data):
    if b'<!DOCTYPE' in data.upper() or b'<!ENTITY' in data.upper():
        raise ValueError('Este documento contiene elementos no admitidos.')
    return ET.fromstring(data)

def _ocr(image):
    from PIL import ImageOps
    import pytesseract
    if image.width*image.height>24_000_000:
        raise ValueError('La imagen es demasiado grande. Usa una copia de menor tamaño.')
    image=ImageOps.exif_transpose(image).convert('RGB')
    image.thumbnail((2400,2400))
    result=pytesseract.image_to_data(ImageOps.autocontrast(ImageOps.grayscale(image)),lang='spa',config='--psm 3',timeout=20,output_type=pytesseract.Output.DICT)
    lines={}; reliable=0
    for i,word in enumerate(result.get('text',[])):
        if not str(word).strip():continue
        try:conf=float(result['conf'][i])
        except (ValueError,TypeError,KeyError,IndexError):conf=0
        if conf>=50:reliable+=1
        line=tuple(result[k][i] for k in ('block_num','par_num','line_num'))
        lines.setdefault(line,[]).append(str(word) if conf>=50 else '[ilegible]')
    return '\n'.join(' '.join(words) for words in lines.values()) if reliable>=2 else ''

def leer_material(nombre,data,pagina_inicio=1,pagina_fin=None):
    from services.planeacion_estabilidad import lecturas_materiales
    cupo=lecturas_materiales()
    if not cupo.acquire(blocking=False):
        raise ValueError('Hay otras transcripciones en proceso. Conserva el archivo cargado y vuelve a pulsar Transcribir en un momento; tu planeación no cambia.')
    try:
        return _leer_material(nombre,data,pagina_inicio,pagina_fin)
    finally:
        cupo.release()


def _leer_material(nombre,data,pagina_inicio=1,pagina_fin=None):
    if not data or len(data)>MAX_BYTES:raise ValueError('Usa archivos de hasta 32 MB.')
    if pagina_inicio<1 or (pagina_fin is not None and pagina_fin<pagina_inicio):raise ValueError('Revisa el intervalo de páginas.')
    extension=Path(nombre).suffix.lower().lstrip('.')
    if extension not in TIPOS:raise ValueError('Convierte este archivo a PDF, Word (.docx), texto o imagen.')
    warnings=[];parts=[]
    try:
        if extension in {'txt','md','csv','json'}:
            for encoding in ('utf-8-sig','utf-16','cp1252'):
                try:parts=[data.decode(encoding)];break
                except UnicodeError:continue
        elif extension=='rtf':
            text=data.decode('cp1252')
            text=re.sub(r"\\'([0-9a-fA-F]{2})",lambda m:bytes.fromhex(m[1]).decode('cp1252'),text)
            text=re.sub(r'\\u(-?\d+)\??',lambda m:chr(int(m[1])%65536),text)
            text=re.sub(r'\\(?:par|line)\b','\n',text)
            parts=[re.sub(r'\\[a-zA-Z]+-?\d* ?|[{}]','',text)]
            warnings.append('Revisa el orden del texto convertido desde RTF.')
        elif extension in {'docx','odt','pptx','odp'}:
            with _zip(data) as archive:
                names=(['word/document.xml'] if extension=='docx' else ['content.xml'] if extension in {'odt','odp'} else
                       sorted([n for n in archive.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',n)],key=lambda n:int(re.search(r'slide(\d+)',n)[1])))
                for name in names:
                    root=_xml(archive.read(name))
                    if extension in {'docx','pptx'}:
                        paragraphs=[e for e in root.iter() if e.tag.endswith('}p')]
                        parts.extend(''.join(e.text or '' for e in p.iter() if e.tag.endswith('}t')) for p in paragraphs)
                    else:
                        parts.extend(''.join(p.itertext()) for p in root.iter() if p.tag.endswith(('}p','}h')))
            warnings.append('Las imágenes incrustadas no se transcriben; si son esenciales, cárgalas por separado.')
        elif extension=='xlsx':
            import openpyxl
            with _zip(data):pass
            wb=openpyxl.load_workbook(BytesIO(data),read_only=True,data_only=True)
            try:
                for sheet in wb.worksheets[:10]:
                    parts.append('Hoja: '+sheet.title)
                    parts.extend(' | '.join(str(v) if v is not None else '' for v in row) for row in sheet.iter_rows(max_row=200,max_col=30,values_only=True) if any(v is not None for v in row))
                    if sheet.max_row>200 or sheet.max_column>30:warnings.append('Se leyó hasta 200 filas y 30 columnas por hoja. Divide tablas más amplias.')
                if len(wb.worksheets)>10:warnings.append('Se leyeron las primeras diez hojas.')
            finally:wb.close()
        elif extension=='pdf':
            import pymupdf
            from PIL import Image
            with pymupdf.open(stream=data,filetype='pdf') as pdf:
                if pdf.needs_pass:raise ValueError('Quita la contraseña del PDF antes de cargarlo.')
                if len(pdf)>MAX_PAGINAS:raise ValueError('El PDF admite hasta 1 000 páginas. Selecciona o divide documentos mayores.')
                if pagina_inicio>len(pdf):raise ValueError('La página inicial no existe en este PDF.')
                fin=min(pagina_fin or len(pdf),len(pdf));ocr_usados=0;omitidas=[];caracteres=0
                for i in range(pagina_inicio-1,fin):
                    page=pdf[i]
                    text=page.get_text().strip()
                    if len(re.sub(r'\W','',text))<30:
                        if ocr_usados>=MAX_OCR:
                            omitidas.append(i+1);continue
                        ocr_usados+=1
                        pix=page.get_pixmap(matrix=pymupdf.Matrix(1.5,1.5),alpha=False)
                        try:text=_ocr(Image.open(BytesIO(pix.tobytes('png'))))
                        except Exception:
                            text='';warnings.append(f'No se pudo transcribir la página {i+1}; el resto se conserva.')
                    parts.append(f'Página {i+1}\n'+(text or '[página sin texto legible]'))
                    caracteres+=len(parts[-1])
                    if caracteres>=MAX_TEXTO:
                        warnings.append(f'Se detuvo la lectura en la página {i+1} al alcanzar el límite de texto. Elige otro intervalo para continuar.');break
                if omitidas:warnings.append(f'Quedan {len(omitidas)} páginas escaneadas sin transcribir, desde la {omitidas[0]}. Selecciona ese intervalo: se procesan hasta diez páginas escaneadas por lectura.')
        else:
            from PIL import Image
            with Image.open(BytesIO(data)) as image:
                frames=getattr(image,'n_frames',1)
                if frames>MAX_PAGINAS:raise ValueError('La imagen multipágina admite hasta 1 000 páginas.')
                if pagina_inicio>frames:raise ValueError('La página inicial no existe en esta imagen.')
                fin=min(pagina_fin or frames,frames,pagina_inicio-1+MAX_OCR)
                for i in range(pagina_inicio-1,fin):
                    image.seek(i);parts.append(_ocr(image.copy()))
                if fin<frames:warnings.append(f'Lectura hasta la página {fin}. Selecciona otro intervalo para continuar.')
            warnings.append('Revisa la transcripción, especialmente la letra manuscrita y las marcas [ilegible].')
    except (ValueError,RuntimeError):raise
    except Exception as exc:
        raise ValueError('No se pudo leer este archivo. Usa una copia en PDF, texto o imagen clara.') from exc
    text='\n'.join(p for p in parts if p.strip()).strip()
    if not text or not re.sub(r'\[página sin texto legible\]|\[ilegible\]|Página \d+','',text).strip():
        raise ValueError('No hay texto legible para incorporar. Prueba otra copia o escribe el contenido.')
    if len(text)>MAX_TEXTO:
        text=text[:MAX_TEXTO];warnings.append('Lectura parcial: hasta 100 000 caracteres. Selecciona las páginas o el extracto pertinente.')
    return {'id':hashlib.sha256(data).hexdigest(),'nombre':Path(nombre).name[:160], 'texto':text,'avisos':warnings}

def incorporar(doc,material,texto):
    from copy import deepcopy
    texto=str(texto).strip()
    if not texto or len(texto)>MAX_TEXTO:raise ValueError('Revisa el texto: máximo 100 000 caracteres por material.')
    result=deepcopy(doc)
    alumnos=list(dict.fromkeys(material.get('alumnos',[])))
    autorizados={a['ID_Alumno'] for a in doc['datos']['Alumnos']}
    if not set(alumnos)<=autorizados:raise PermissionError('Este material incluye alumnos fuera de la planeación.')
    materiales=[m for m in result['metadatos'].get('materiales',[]) if m['id']!=material['id']]
    if len(materiales)>=10:raise ValueError('Puedes conservar hasta diez materiales por planeación. Resume o reemplaza uno.')
    if sum(len(m['texto']) for m in materiales)+len(texto)>MAX_TOTAL_TEXTO:
        raise ValueError('El borrador admite 120 000 caracteres de materiales en total. Selecciona los fragmentos útiles; los originales no se modifican.')
    materiales.append({'id':material['id'],'nombre':material['nombre'],'texto':texto,'revisado':True,'alumnos':alumnos})
    result['metadatos']['materiales']=materiales
    return result

def fuentes_materiales(doc):
    return [{'tipo':'MATERIAL','registro':m['id'],'referencia':'M'+str(i+1),'fecha':'',
             'alumno':','.join(m.get('alumnos',[])),'vinculo_verificado':bool(m.get('alumnos')),
             'texto':m['texto'][:6000],'alcance':'extracto del material revisado por docente (hasta 6 000 caracteres); el texto incorporado completo permanece en Materiales; no resultado ni fuente oficial automática'}
            for i,m in enumerate(doc['metadatos'].get('materiales',[])) if m.get('revisado') and m.get('texto')]
