import unittest
from io import BytesIO
from unittest.mock import patch
from copy import deepcopy
from docx import Document
from services.materiales_planeacion import leer_material, incorporar, fuentes_materiales, MAX_BYTES, MAX_TEXTO
from services.planeacion_modelo import crear_plantilla
from services import planeacion

def documento():
    alumno={'ID_Alumno':'A1','Nombre_Completo':'Alumno ficticio','ID_Escuela':'E1','Nombre_Escuela':'Escuela ficticia','Grado':'5','Grupo':'A'}
    doc=crear_plantilla('XXIII',['A1'],[alumno],autor='Autor ficticio',funcion='Aprendizaje')
    doc.update(id='prueba',cuenta='prueba')
    return doc

class MaterialesTests(unittest.TestCase):
    def test_word_incluye_parrafos_y_tablas_sin_ejecutar(self):
        doc=Document();doc.add_paragraph('Lectura con apoyo visual')
        table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='Necesidad';table.cell(0,1).text='Instrucciones breves'
        buffer=BytesIO();doc.save(buffer)
        result=leer_material('guia.docx',buffer.getvalue())
        self.assertIn('Lectura con apoyo visual',result['texto']);self.assertIn('Instrucciones breves',result['texto'])
        self.assertIn('incrustadas',result['avisos'][0])
    def test_pdf_texto_no_necesita_ocr(self):
        import pymupdf
        with pymupdf.open() as doc:
            doc.new_page().insert_text((30,30),'Actividad de lectura con apoyo visual e instrucciones breves para trabajar.')
            data=doc.tobytes()
        with patch('services.materiales_planeacion._ocr',side_effect=AssertionError('No debe llamar OCR')):
            self.assertIn('Actividad de lectura',leer_material('material.pdf',data)['texto'])
    def test_imagen_local_y_pdf_escaneado(self):
        from PIL import Image
        image=Image.new('RGB',(100,100),'white');buffer=BytesIO();image.save(buffer,format='PNG')
        with patch('services.materiales_planeacion._ocr',return_value='Actividad transcrita [ilegible]'):
            result=leer_material('foto.png',buffer.getvalue())
        self.assertIn('[ilegible]',result['texto'])
    def test_sin_texto_limites_y_formato_invalido(self):
        for nombre,data in [('material.exe',b'no ejecutar'),('grande.txt',b'a'*(MAX_BYTES+1)),('vacio.txt',b'')]:
            with self.assertRaises(ValueError):leer_material(nombre,data)
        self.assertEqual(len(leer_material('largo.txt',b'x'*50000)['texto']),50000)
        self.assertIn('parcial',leer_material('largo.txt',b'x'*(MAX_TEXTO+1))['avisos'][0])
    def test_pdf_largo_y_lectura_por_intervalo(self):
        import pymupdf
        with pymupdf.open() as doc:
            for i in range(150):doc.new_page().insert_text((30,30),f'Pagina {i+1}: actividad educativa con instrucciones claras y apoyos visuales.')
            data=doc.tobytes()
        result=leer_material('guia.pdf',data)
        self.assertIn('Página 150',result['texto'])
        parcial=leer_material('guia.pdf',data,140,150)
        self.assertIn('Página 140',parcial['texto']);self.assertNotIn('Página 139',parcial['texto'])
    def test_pdf_escaneado_limita_ocr_y_avisa(self):
        import pymupdf
        with pymupdf.open() as doc:
            for _ in range(14):doc.new_page()
            data=doc.tobytes()
        with patch('services.materiales_planeacion._ocr',return_value='Texto revisable para trabajar') as ocr:
            result=leer_material('escaneo.pdf',data)
        self.assertEqual(ocr.call_count,10)
        self.assertTrue(any('sin transcribir' in w for w in result['avisos']))
    def test_presupuesto_de_materiales_protege_respaldo(self):
        d=documento();material=leer_material('uno.txt',b'x'*80000);d=incorporar(d,material,material['texto'])
        otro=leer_material('dos.txt',b'y'*80000)
        with self.assertRaisesRegex(ValueError,'120 000'):incorporar(d,otro,otro['texto'])
    def test_material_revisado_persistente_idempotente_y_contexto(self):
        source=documento();material=leer_material('guia.txt','Apoyo visual documentado'.encode())
        doc=incorporar(source,material,'Texto revisado por docente')
        doc=incorporar(doc,material,'Texto corregido nuevamente')
        self.assertEqual(len(doc['metadatos']['materiales']),1)
        self.assertNotIn('materiales',source['metadatos'])
        with patch.object(planeacion,'autorizar'),patch.object(planeacion,'evidencias',return_value=([],[])):
            actualizado=planeacion.preparar_contexto(doc)
        self.assertIn('Texto corregido nuevamente',actualizado['metadatos']['resumen_educativo'])
        self.assertEqual(actualizado['metadatos']['fuentes'][0]['tipo'],'MATERIAL')
        self.assertNotIn('guia.txt',actualizado['metadatos']['resumen_educativo'])
    def test_varios_formatos_y_zip_hostil(self):
        import zipfile
        for extension in ('md','csv','json'):
            self.assertTrue(leer_material('material.'+extension,b'Contenido revisable')['texto'])
        buffer=BytesIO()
        with zipfile.ZipFile(buffer,'w') as archive:
            archive.writestr('content.xml','<!DOCTYPE x [<!ENTITY peligro "texto">]><x>&peligro;</x>')
        with self.assertRaises(ValueError):leer_material('archivo.odt',buffer.getvalue())
    def test_formatos_office_abiertos_y_presentacion(self):
        import zipfile
        buffer=BytesIO()
        with zipfile.ZipFile(buffer,'w') as archive:
            archive.writestr('content.xml','<document xmlns:t="urn:text"><t:p>Actividad de apoyo</t:p></document>')
        for ext in ('odt','odp'):
            self.assertEqual(leer_material('archivo.'+ext,buffer.getvalue())['texto'],'Actividad de apoyo')
        buffer=BytesIO()
        with zipfile.ZipFile(buffer,'w') as archive:
            archive.writestr('ppt/slides/slide1.xml','<slide xmlns:a="urn:ppt"><a:p><a:r><a:t>Lectura compartida</a:t></a:r></a:p></slide>')
        self.assertEqual(leer_material('archivo.pptx',buffer.getvalue())['texto'],'Lectura compartida')
        self.assertIn('Actividad',leer_material('archivo.rtf',b'{\\rtf1 Actividad\\par Lectura}')['texto'])
    def test_resumen_conserva_materiales_y_adaptaciones_sin_duplicarlos(self):
        from ui.planeacion import _resumen_complementos
        doc=documento();doc['metadatos']['resumen_educativo']='Evidencia previa '
        for i in range(5):
            material=leer_material(f'material{i}.txt',f'MATERIAL_{i} texto pertinente'.encode())
            doc=incorporar(doc,material,material['texto'])
        doc['metadatos']['curriculo']={'adaptaciones':{'R1':{'contenido':'CONTENIDO ADAPTADO','pda':'OBJETIVO ADAPTADO'}}}
        with patch.object(planeacion,'autorizar'):
            once=_resumen_complementos(deepcopy(doc))
            twice=_resumen_complementos(deepcopy(once))
        resumen=twice['metadatos']['resumen_educativo']
        for i in range(5):self.assertIn(f'MATERIAL_{i}',resumen)
        self.assertIn('OBJETIVO ADAPTADO',resumen)
        self.assertEqual(resumen.count('COMPLEMENTOS REVISADOS:'),1)
        self.assertLessEqual(len(resumen),12000)

if __name__=='__main__':unittest.main()
